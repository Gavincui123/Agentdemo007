#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""多格式文档「清洗-切分-入库」流水线 —— 客服知识库 RAG 前置（纯 Python，不依赖 LlamaIndex）。

四层架构：
  readers   分型解析 10 类文档（md/txt、html、docx、pdf 文本型、pdf 扫描型、xlsx、csv、
            json 活动规则、纯图片、嵌套表格）→ Block 列表（标题层级/页码/sheet/时间约束 信息）
  cleaners  清洗：NFKC 规范化、PDF 页眉页脚去重、乱码剔除、近似重复块去重、PII 脱敏钩子
  chunkers  分层切分：structure 按标题层级（默认，FAQ 每问一答原子块）
            | semantic 句子嵌入余弦断点（txt/扫描 OCR 等无标题长文本）
            表格恒为原子块不拆散；--parent-child 小块检索大块返回；硬约束 max-chars/overlap 兜底
  writer    嵌入（SiliconFlow Qwen/Qwen3-Embedding-8B，4096 维）→ Chroma / Milvus 幂等 upsert
            → 库内数量核对 + 语义召回自检

用法示例：
  # 1) markdown 语料入库（连远程已部署的 Chroma，端口 8000）
  python ingest.py --input ../../src/main/resources/corpus --store chroma --uri http://<服务器IP>:8000
  # 服务端开启 token 鉴权时：export CHROMA_TOKEN=<token>；https 端点直接写 https:// 即可
  # 2) 混合样例目录入库（先跑 make_samples.py 生成）
  python ingest.py --input ./samples --store milvus --uri http://localhost:19530
  # 3) 只解析切分不落库（调切块策略用）
  python ingest.py --input ../../src/main/resources/corpus --dry-run

与 Java 侧对齐约定（重要，详见 docs/rag-corpus-ingestion-tutorial.md）：
  - 嵌入模型必须与后端 EmbeddingService 同源（Qwen/Qwen3-Embedding-8B，4096 维），换模型 = 换向量空间，全库必须重灌；
  - metadata 契约：source/doc_type/page/sheet/section/doc_no/domain/temporal_tag/kind/parent_text/valid_from/valid_until；
  - temporal_tag=HISTORICAL 的块将被后端时效治理隔离标注；JSON 活动规则的 valid_from/valid_until
    表达「活动时间窗口」（已结束/进行中/未开始三种时间态由此判定），勿把现行政策误标为历史。
"""

from __future__ import annotations

import argparse
import csv
import difflib
import hashlib
import json
import os
import re
import sys
import time
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

# ---------------------------------------------------------------------------
# 数据结构
# ---------------------------------------------------------------------------

KIND_TEXT, KIND_TABLE, KIND_FIGURE, KIND_HEADING = "text", "table", "figure", "heading"


@dataclass
class Block:
    """reader 输出的最小内容单元（一个标题行 / 一段正文 / 一张表 / 一张图 / 一条结构化记录）。"""

    text: str
    kind: str = KIND_TEXT
    level: int = 0            # 标题层级 1-6（非标题为 0）
    page: int | None = None   # PDF 页码（1 起）
    sheet: str | None = None  # Excel 工作表名
    source: str = ""          # 相对文件路径
    doc_type: str = ""        # md/txt/html/docx/pdf/pdf_ocr/xlsx/csv/json/image_ocr/image
    atomic: bool = False      # 原子块：JSON 活动条目等结构化独立单元，不与相邻正文合并
    valid_from: str | None = None    # 时间约束（活动规则类语料：生效日）
    valid_until: str | None = None   # 时间约束（活动规则类语料：截止日）


@dataclass
class DocMeta:
    """文档级元数据（markdown 首行引用块 / 文件属性解析而来）。"""

    doc_no: str | None = None
    domain: str | None = None
    updated: str | None = None
    temporal_tag: str = "CURRENT"
    valid_until: str | None = None


@dataclass
class Chunk:
    """最终落库单元：确定 id + 正文 + 元数据（与 Java 侧 RagFragment 契约对应）。"""

    id: str
    text: str
    meta: dict = field(default_factory=dict)


# ---------------------------------------------------------------------------
# cleaners —— 清洗层
# ---------------------------------------------------------------------------

PHONE_RE = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")
IDCARD_RE = re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)")


def normalize_text(text: str) -> str:
    """NFKC 规范化 + 空白折叠（保留换行结构；表格块不走此函数以免破坏竖线）。"""
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u3000]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def strip_garbled(text: str) -> str | None:
    """乱码剔除：U+FFFD 占比过高的块丢弃（返回 None），少量替换符就地清除。"""
    if not text:
        return None
    if text.count("\ufffd") > max(2, len(text) * 0.05):
        return None
    return text.replace("\ufffd", "")


def redact_pii(text: str) -> str:
    """脱敏钩子：手机号 / 身份证打码。知识库一般无 PII，作为入库前保险丝。"""
    text = PHONE_RE.sub("1**********", text)
    text = IDCARD_RE.sub(lambda m: m.group(0)[:3] + "***********" + m.group(0)[-3:], text)
    return text


def strip_pdf_headers_footers(pages: list[list[str]], min_repeat_ratio: float = 0.6) -> None:
    """PDF 页眉页脚去重：出现在 >=60% 页面的重复行（页眉/页脚/水印行）从每页剔除（原地修改）。

    页数 < 3 时不启用（样本不足以判定重复）。ratio 可调：重复表格头被误删时调高。
    """
    if len(pages) < 3:
        return
    counter: Counter[str] = Counter()
    for lines in pages:
        counter.update({ln.strip() for ln in lines if len(ln.strip()) >= 3})
    threshold = max(3, int(len(pages) * min_repeat_ratio))
    repeated = {ln for ln, n in counter.items() if n >= threshold}
    if not repeated:
        return
    for lines in pages:
        lines[:] = [ln for ln in lines if ln.strip() not in repeated]


def _quick_ratio(a: str, b: str) -> float:
    return difflib.SequenceMatcher(None, a, b).quick_ratio()


def clean_blocks(blocks: list[Block]) -> list[Block]:
    """清洗主入口：规范化 + 乱码剔除 + 脱敏 + 近似重复块去重 + 空块过滤。

    表格块不做空白折叠（保住 markdown 竖线对齐）；标题块不做最小长度过滤（短标题合法）。
    """
    cleaned: list[Block] = []
    seen_exact: set[str] = set()
    for b in blocks:
        if b.kind == KIND_TABLE:
            text = redact_pii(b.text)
            if not text or "\ufffd" in text:
                continue
        else:
            text = strip_garbled(normalize_text(b.text))
            if text is None:
                continue
            text = redact_pii(text)
            if b.kind != KIND_HEADING and len(re.sub(r"\s", "", text)) < 4:
                continue  # 有效内容过短的碎块丢弃（短标题除外）
        key = re.sub(r"\s", "", text)
        if not key:
            continue
        if key in seen_exact:
            continue
        if b.kind == KIND_TEXT and len(key) <= 200:
            # 近似去重：只与最近 50 个文本块比较，KB 规模 O(n·50) 可接受
            if any(_quick_ratio(key, _norm_prev) > 0.95 for _norm_prev in _recent_norm):
                continue
        seen_exact.add(key)
        if b.kind == KIND_TEXT:
            _recent_norm.append(key)
            del _recent_norm[:-50]
        cleaned.append(Block(text=text, kind=b.kind, level=b.level, page=b.page,
                             sheet=b.sheet, source=b.source, doc_type=b.doc_type,
                             atomic=b.atomic, valid_from=b.valid_from,
                             valid_until=b.valid_until))
    return cleaned


_recent_norm: list[str] = []  # clean_blocks 内部的近重比较窗口（每次调用前应清空）


def clean_document(blocks: list[Block]) -> list[Block]:
    """单文档清洗入口（重置近重窗口后调 clean_blocks）。"""
    _recent_norm.clear()
    return clean_blocks(blocks)


# ---------------------------------------------------------------------------
# readers —— 分型解析层（10 类）
# ---------------------------------------------------------------------------

META_LINE_RE = {
    "doc_no": re.compile(r"文档编号[:：]\s*([A-Za-z0-9\-]+)"),
    "domain": re.compile(r"知识域[:：]\s*([a-z_]+)"),
    "updated": re.compile(r"更新日期[:：]\s*(\d{4}-\d{2}-\d{2})"),
    "temporal_tag": re.compile(r"时效标签[:：]\s*(CURRENT|HISTORICAL)"),
    "valid_until": re.compile(r"有效期至[:：]\s*(\d{4}-\d{2}-\d{2})"),
}


def parse_meta_line(text: str) -> DocMeta:
    """解析元数据行：『> 文档编号：KB-X ｜ 知识域：yyy ｜ 更新日期：… ｜ 时效标签：CURRENT』"""
    meta = DocMeta()
    for field_name, pattern in META_LINE_RE.items():
        m = pattern.search(text)
        if m:
            setattr(meta, field_name, m.group(1))
    return meta


def _parse_markdown_raw(raw: str, filename: str, doc_type: str) -> tuple[list[Block], DocMeta]:
    """统一 markdown 解析：`#` 标题建层级；`|` 表格行聚合为原子表块；首部引用块解析元数据。

    html（BeautifulSoup 转出）与 docx（python-docx 转出）的中间 markdown 均复用此函数。
    """
    meta = DocMeta()
    lines = raw.splitlines()
    meta_line_idx: set[int] = set()
    for i, ln in enumerate(lines[:12]):
        # 元数据行指纹：同时含「文档编号」与「知识域」（md 引用块 / docx 段落两种来源通用）
        if "文档编号" in ln and "知识域" in ln:
            meta = parse_meta_line(ln)
            meta_line_idx.add(i)  # 机读元数据行已解析入库，不再作为正文切块
            break

    blocks: list[Block] = []
    table_buf: list[str] = []
    text_buf: list[str] = []

    def flush(buf: list[str], kind: str) -> None:
        if not buf:
            return
        joined = "\n".join(buf).strip()
        if joined:
            blocks.append(Block(text=joined, kind=kind, source=filename, doc_type=doc_type))
        buf.clear()

    for i, ln in enumerate(lines):
        if i in meta_line_idx:
            continue
        stripped = ln.strip()
        if stripped.startswith("|"):
            flush(text_buf, KIND_TEXT)
            table_buf.append(stripped)
        elif stripped.startswith("#"):
            flush(text_buf, KIND_TEXT)
            flush(table_buf, KIND_TABLE)
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append(Block(text=stripped.lstrip("#").strip(), kind=KIND_HEADING,
                                level=level, source=filename, doc_type=doc_type))
        else:
            flush(table_buf, KIND_TABLE)
            text_buf.append(ln)
    flush(text_buf, KIND_TEXT)
    flush(table_buf, KIND_TABLE)
    return blocks, meta


def read_markdown(path: Path) -> tuple[list[Block], DocMeta]:
    """md：直接读（LLM 友好，保留标题结构）。"""
    return _parse_markdown_raw(path.read_text(encoding="utf-8"), path.name, "md")


def read_txt(path: Path) -> tuple[list[Block], DocMeta]:
    """txt：整文件一个文本块（无标题结构，交由 semantic 切分或硬约束兜底）。"""
    text = path.read_text(encoding="utf-8")
    return [Block(text=text, kind=KIND_TEXT, source=path.name, doc_type="txt")], DocMeta()


def read_html(path: Path) -> tuple[list[Block], DocMeta]:
    """html：BeautifulSoup 去 script/style/nav，h1-h6 转标题、表格转 md 表，复用 md 解析。"""
    try:
        from bs4 import BeautifulSoup
    except ImportError as e:
        raise SystemExit("缺少依赖：pip install beautifulsoup4") from e
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    for tag in soup(["script", "style", "nav", "header", "footer", "aside"]):
        tag.decompose()
    md_lines: list[str] = []
    for el in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "tr"]):
        text = el.get_text(" ", strip=True)
        if not text:
            continue
        if el.name.startswith("h"):
            md_lines.append("#" * int(el.name[1]) + " " + text)
        elif el.name == "tr":
            cells = el.find_all(["th", "td"])
            md_lines.append("| " + " | ".join(c.get_text(" ", strip=True) for c in cells) + " |")
        else:
            md_lines.append(text)
    return _parse_markdown_raw("\n".join(md_lines), path.name, "html")


def read_docx(path: Path) -> tuple[list[Block], DocMeta]:
    """docx：python-docx 按文档顺序遍历段落与表格；标题样式映射层级；表格序列化 md 原子块。"""
    try:
        import docx  # python-docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as e:
        raise SystemExit("缺少依赖：pip install python-docx") from e
    document = docx.Document(str(path))
    meta = DocMeta()
    md_lines: list[str] = []

    def heading_level(paragraph) -> int:
        name = (paragraph.style.name or "") if paragraph.style is not None else ""
        m = re.match(r"(?:Heading|标题)\s*(\d)", name, re.IGNORECASE)
        return int(m.group(1)) if m else 0

    for element in document.element.body.iterchildren():
        tag = element.tag.split("}")[-1]
        if tag == "p":
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            if not text:
                continue
            level = heading_level(paragraph)
            md_lines.append(("#" * level + " " + text) if level else text)
        elif tag == "tbl":
            table = Table(element, document)
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                md_lines.append("| " + " | ".join(cells) + " |")
            md_lines.append("")  # 表格后断行，避免与后续段落粘连
    blocks, meta = _parse_markdown_raw("\n".join(md_lines), path.name, "docx")
    try:
        core = document.core_properties
        if core.modified is not None and meta.updated is None:
            meta.updated = core.modified.strftime("%Y-%m-%d")
    except Exception:
        pass
    return blocks, meta


def _table_to_md(rows: list[list[object]]) -> str:
    """二维行转 markdown 表（None→空串，竖线转义，空行剔除）。"""

    def cell(v) -> str:
        return str(v if v is not None else "").replace("|", "\\|").replace("\n", " ").strip()

    rows = [[cell(c) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    out.extend("| " + " | ".join(r) + " |" for r in rows[1:])
    return "\n".join(out)


def _rows_to_table_blocks(header: list[object], body: list[list[object]],
                          sheet: str, source: str, doc_type: str,
                          max_rows: int = 50) -> list[Block]:
    """表格落块策略：小表整表一块；大表按 max_rows 行分块、每块重复表头（保持列语义）。"""
    if len(body) <= max_rows:
        return [Block(text=_table_to_md(header + body), kind=KIND_TABLE, sheet=sheet,
                      source=source, doc_type=doc_type)]
    blocks: list[Block] = []
    for i in range(0, len(body), max_rows):
        part = body[i : i + max_rows]
        note = f"\n（{sheet or '表'} 第 {i + 1}-{i + len(part)} 行，共 {len(body)} 行）"
        blocks.append(Block(text=_table_to_md([header] + part) + note, kind=KIND_TABLE,
                            sheet=sheet, source=source, doc_type=doc_type))
    return blocks


def read_xlsx(path: Path) -> tuple[list[Block], DocMeta]:
    """xlsx：每 sheet 一个逻辑表（表头拼进行内保持列语义），大 sheet 分块带头。"""
    try:
        import openpyxl
    except ImportError as e:
        raise SystemExit("缺少依赖：pip install openpyxl") from e
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    blocks: list[Block] = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(v is not None and str(v).strip() for v in r)]
        if not rows:
            continue
        blocks.extend(_rows_to_table_blocks(rows[0], rows[1:], ws.title, path.name, "xlsx"))
    return blocks, DocMeta()


def read_csv(path: Path) -> tuple[list[Block], DocMeta]:
    """csv/tsv：与 xlsx 同策略。"""
    sep = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open(encoding="utf-8-sig", newline="") as f:
        rows = [r for r in csv.reader(f, delimiter=sep) if any(c.strip() for c in r)]
    if not rows:
        return [], DocMeta()
    return _rows_to_table_blocks(rows[0], rows[1:], None, path.name, "csv"), DocMeta()


ACT_LABELS = {
    "type": "活动类型",
    "rules": "活动规则",
    "restrictions": "限制说明",
    "superposition": "叠加规则",
    "member_levels": "适用会员",
    "channels": "适用渠道",
    "notes": "备注",
}


def _activity_to_text(item: dict) -> str:
    """JSON 活动条目 → 可检索文本（键名中文化、列表分号拼接），未识别字段兜底输出。"""
    lines = [f"【活动规则】{item.get('name', item.get('id', '活动'))}"
             + (f"（{item['type']}）" if item.get("type") else "")]
    t = item.get("time") or {}
    start, end = t.get("start"), t.get("end")
    lines.append("活动时间：" + (f"{start} 至 {end}" if start or end else str(t or "长期")))
    for key, label in ACT_LABELS.items():
        value = item.get(key)
        if value is None or key == "type":
            continue
        if isinstance(value, list):
            value = "；".join(str(v) for v in value)
        lines.append(f"{label}：{value}")
    for key, value in item.items():  # 未知字段兜底（不丢信息）
        if key in ACT_LABELS or key in ("name", "id", "time"):
            continue
        if isinstance(value, list):
            value = "；".join(str(v) for v in value)
        lines.append(f"{key}：{value}")
    return "\n".join(lines)


def read_json(path: Path) -> tuple[list[Block], DocMeta]:
    """json（活动规则类）：每条记录一个原子块，time.start/end 提取为 valid_from/valid_until 元数据。

    支持 {"doc_no","domain","updated","temporal_tag","activities":[...]} 或裸数组 [...]。
    活动时间约束是时效测试的关键：已结束/进行中/未开始的条目靠 valid_from/valid_until 区分，
    Java 侧检索后可据此判断"该活动现在还能不能参加"。
    """
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, dict):
        meta = DocMeta(doc_no=data.get("doc_no"), domain=data.get("domain"),
                       updated=data.get("updated"),
                       temporal_tag=data.get("temporal_tag", "CURRENT") or "CURRENT")
        items = data.get("activities") or data.get("items") or []
    elif isinstance(data, list):
        meta, items = DocMeta(), data
    else:
        return [], DocMeta()
    blocks = [
        Block(text=_activity_to_text(item), kind=KIND_TEXT, atomic=True,
              source=path.name, doc_type="json",
              valid_from=(item.get("time") or {}).get("start"),
              valid_until=(item.get("time") or {}).get("end"))
        for item in items if isinstance(item, dict)
    ]
    return blocks, meta


_OCR_ENGINE = None  # PaddleOCR 引擎进程内缓存（首次加载很慢）


def _ocr_available() -> bool:
    try:
        import paddleocr  # noqa: F401
        return True
    except ImportError:
        return False


def _ocr_image(pil_image) -> str:
    """对 PIL 图片跑 PaddleOCR（中文），返回按行拼接文本；失败抛异常由调用方降级。"""
    global _OCR_ENGINE
    import numpy as np
    from paddleocr import PaddleOCR

    if _OCR_ENGINE is None:
        _OCR_ENGINE = PaddleOCR(use_angle_cls=True, lang="ch", show_log=False)
    result = _OCR_ENGINE.ocr(np.asarray(pil_image), cls=True)
    lines: list[str] = []
    for page in result or []:
        for item in page or []:
            lines.append(item[1][0])
    return "\n".join(lines)


def read_pdf(path: Path, ocr_threshold_chars: int = 30) -> tuple[list[Block], DocMeta]:
    """pdf：文本型走 PyMuPDF 版面块 + 加粗/短行识别标题；pdfplumber 抽表格；低文本页判扫描页走 OCR。

    扫描页判定：单页有效文本 < ocr_threshold_chars → 视为扫描页。
      - 安装了 PaddleOCR：整页渲染 200dpi OCR（doc_type=pdf_ocr）
      - 未安装：输出 figure 占位块并告警（不空过、不崩）
    """
    try:
        import fitz  # PyMuPDF
    except ImportError as e:
        raise SystemExit("缺少依赖：pip install PyMuPDF") from e

    doc = fitz.open(str(path))
    page_lines: list[list[str]] = []   # 每页的行文本（供页眉页脚判定）
    page_span_blocks: list[list[Block]] = []  # 每页的行级块
    for pno in range(len(doc)):
        page = doc.load_page(pno)
        d = page.get_text("dict")
        lines: list[str] = []
        blocks: list[Block] = []
        for blk in d.get("blocks", []):
            if blk.get("type") != 0:  # 非文本块（图像块另行统计）
                continue
            for line in blk.get("lines", []):
                spans = line.get("spans", [])
                if not spans:
                    continue
                text = "".join(s.get("text", "") for s in spans).strip()
                if not text:
                    continue
                span = spans[0]
                bold = bool(span.get("flags", 0) & 16)
                lines.append(text)
                blocks.append(Block(text=text,
                                    kind=KIND_HEADING if (bold and len(text) < 40) else KIND_TEXT,
                                    page=pno + 1, source=path.name, doc_type="pdf"))
        page_lines.append(lines)
        page_span_blocks.append(blocks)

    strip_pdf_headers_footers(page_lines)
    kept_lines = [set(lines) for lines in page_lines]  # 剔除页眉页脚后每页保留的行

    can_ocr = _ocr_available()
    blocks: list[Block] = []
    for pno in range(len(doc)):
        if sum(len(t) for t in page_lines[pno]) < ocr_threshold_chars:  # 扫描页
            if can_ocr:
                pix = doc.load_page(pno).get_pixmap(dpi=200)
                import io

                from PIL import Image
                ocr_text = _ocr_image(Image.open(io.BytesIO(pix.tobytes("png"))))
                if ocr_text.strip():
                    blocks.append(Block(text=ocr_text, kind=KIND_TEXT, page=pno + 1,
                                        source=path.name, doc_type="pdf_ocr"))
                    continue
            print(f"[warn] {path.name} 第 {pno + 1} 页疑似扫描页且 OCR 不可用，输出占位块", file=sys.stderr)
            blocks.append(Block(text=f"【扫描页】第 {pno + 1} 页为图像页，未能提取文字，请人工补录",
                                kind=KIND_FIGURE, page=pno + 1, source=path.name, doc_type="pdf"))
            continue
        for b in page_span_blocks[pno]:
            if b.text not in kept_lines[pno]:
                continue  # 该行被判定为页眉页脚，剔除
            blocks.append(b)
        images = doc.load_page(pno).get_images(full=True)
        if images:
            blocks.append(Block(text=f"【图片】第 {pno + 1} 页含 {len(images)} 张图像（图注缺失，请人工核对）",
                                kind=KIND_FIGURE, page=pno + 1, source=path.name, doc_type="pdf"))
    doc.close()

    try:
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            for pno, page in enumerate(pdf.pages, start=1):
                for table in page.extract_tables() or []:
                    md = _table_to_md(table)
                    if md:
                        blocks.append(Block(text=md, kind=KIND_TABLE, page=pno,
                                            source=path.name, doc_type="pdf"))
    except ImportError:
        print("[warn] 未装 pdfplumber，跳过 PDF 表格抽取（pip install pdfplumber）", file=sys.stderr)
    return blocks, DocMeta()


def read_image(path: Path) -> tuple[list[Block], DocMeta]:
    """纯图片：PaddleOCR 可用则整图 OCR 为文本块；否则输出占位元数据块（不空过、不崩）。"""
    try:
        from PIL import Image
    except ImportError as e:
        raise SystemExit("缺少依赖：pip install Pillow") from e
    if _ocr_available():
        text = _ocr_image(Image.open(str(path)).convert("RGB"))
        if text.strip():
            return [Block(text=text, kind=KIND_TEXT, source=path.name, doc_type="image_ocr")], DocMeta()
    return [Block(text=f"【图片】{path.name}：OCR 不可用或图内无可识别文字，请人工补录文字说明",
                  kind=KIND_FIGURE, source=path.name, doc_type="image")], DocMeta()


READERS = {
    ".md": read_markdown,
    ".markdown": read_markdown,
    ".txt": read_txt,
    ".html": read_html,
    ".htm": read_html,
    ".docx": read_docx,
    ".pdf": read_pdf,
    ".xlsx": read_xlsx,
    ".xls": read_xlsx,
    ".csv": read_csv,
    ".tsv": read_csv,
    ".json": read_json,
    ".png": read_image,
    ".jpg": read_image,
    ".jpeg": read_image,
}

SUPPORTED_EXTS = sorted(READERS.keys())

DOC_TYPE_LABEL = {"markdown": "md", "htm": "html", "tsv": "csv",
                  "jpg": "image", "jpeg": "image", "png": "image", "xls": "xlsx"}


# ---------------------------------------------------------------------------
# chunkers —— 切分层
# ---------------------------------------------------------------------------

SENT_SPLIT_RE = re.compile(r"(?<=[。！？；!?;])\s*")


def split_sentences(text: str) -> list[str]:
    """中文句号类标点 + 换行切句（保留标点）。"""
    sents: list[str] = []
    for seg in text.split("\n"):
        seg = seg.strip()
        if not seg:
            continue
        parts = [p for p in SENT_SPLIT_RE.split(seg) if p.strip()]
        sents.extend(parts or [seg])
    return sents


def hard_cap(text: str, max_chars: int, overlap: int) -> list[str]:
    """硬约束兜底：句子边界贪心装填，超 max-chars 切块、overlap 字符回滚重叠。"""
    if len(text) <= max_chars:
        return [text]
    pieces: list[str] = []
    cur = ""
    for s in split_sentences(text):
        if cur and len(cur) + len(s) + 1 > max_chars:
            pieces.append(cur)
            cur = (cur[-overlap:] + s) if overlap > 0 else s
        else:
            cur = (cur + s) if cur else s
        while len(cur) > max_chars:  # 单句超长的极端情况
            pieces.append(cur[:max_chars])
            cur = cur[max_chars - overlap:]
    if cur.strip():
        pieces.append(cur)
    return pieces


def _mk_chunk(text: str, meta: DocMeta, source: str, doc_type: str, section: str,
              kind: str, page: int | None = None, sheet: str | None = None,
              valid_from: str | None = None, valid_until: str | None = None) -> Chunk:
    return Chunk(
        id="",
        text=text,
        meta={
            "source": source,
            "doc_type": doc_type,
            "kind": kind,
            "section": section[:180],
            "doc_no": meta.doc_no or "",
            "domain": meta.domain or "",
            "temporal_tag": meta.temporal_tag,
            "updated": meta.updated or "",
            **({"page": page} if page is not None else {}),
            **({"sheet": sheet} if sheet is not None else {}),
            **({"valid_from": valid_from} if valid_from else {}),
            **({"valid_until": valid_until} if valid_until else {}),
        },
    )


def structure_chunks(blocks: list[Block], meta: DocMeta, source: str, doc_type: str,
                     max_chars: int, overlap: int) -> list[Chunk]:
    """structure 切分：标题层级开新节、正文累积；表格恒为原子块（附标题上下文）。

    节文本带上标题路径前缀再嵌入（检索命中率更高）；超 max-chars 的节由硬约束兜底再切。
    """
    chunks: list[Chunk] = []
    path_stack: list[tuple[int, str]] = []
    buf: list[str] = []

    def section_title() -> str:
        return " / ".join(t for _, t in path_stack) if path_stack else "(正文)"

    def flush() -> None:
        if not buf:
            return
        body = "\n".join(buf).strip()
        buf.clear()
        prefix = section_title()
        full = f"{prefix}\n{body}" if prefix != "(正文)" else body
        for piece in hard_cap(full, max_chars, overlap):
            chunks.append(_mk_chunk(piece, meta, source, doc_type,
                                    section=prefix, kind=KIND_TEXT))

    for b in blocks:
        if b.kind == KIND_HEADING:
            flush()
            while path_stack and path_stack[-1][0] >= b.level:
                path_stack.pop()
            path_stack.append((b.level, b.text))
        elif b.kind == KIND_TABLE or b.atomic:
            # 表格与结构化原子块：永不与相邻正文合并，整块独立入库（附标题上下文）
            flush()
            content = f"{section_title()}\n{b.text}" if path_stack else b.text
            for piece in hard_cap(content, max(max_chars, 1200), overlap):
                chunks.append(_mk_chunk(piece, meta, source, doc_type,
                                        section=section_title(), kind=b.kind,
                                        page=b.page, sheet=b.sheet,
                                        valid_from=b.valid_from, valid_until=b.valid_until))
        else:
            buf.append(b.text)
    flush()
    return chunks


def _cosine(a: list[float], b: list[float]) -> float:
    import math

    dot = sum(x * y for x, y in zip(a, b))
    na = math.sqrt(sum(x * x for x in a)) or 1e-9
    nb = math.sqrt(sum(y * y for y in b)) or 1e-9
    return dot / (na * nb)


def semantic_chunks(blocks: list[Block], meta: DocMeta, source: str, doc_type: str,
                    embedder: "Embedder", threshold: float,
                    max_chars: int, overlap: int) -> list[Chunk]:
    """semantic 切分：句子嵌入余弦断点，适合无标题长文本（txt/扫描 OCR 输出）。

    相邻句相似度 < threshold 即断开（threshold 越大切得越碎）；表格块不走此路，仍原子。
    """
    text_blocks = [b for b in blocks if b.kind in (KIND_TEXT, KIND_HEADING)]
    if not text_blocks:
        return structure_chunks(blocks, meta, source, doc_type, max_chars, overlap)
    sentences: list[str] = []
    for b in text_blocks:
        sentences.extend(split_sentences(b.text))
    vectors = embedder.embed(sentences)
    chunks: list[Chunk] = []
    group: list[str] = []
    for i, sent in enumerate(sentences):
        group.append(sent)
        if (i < len(sentences) - 1 and _cosine(vectors[i], vectors[i + 1]) < threshold
                and sum(len(s) for s in group) >= 60):
            chunks.append(_mk_chunk(" ".join(group), meta, source, doc_type,
                                    section="(语义段)", kind=KIND_TEXT))
            group = []
    if group:
        chunks.append(_mk_chunk(" ".join(group), meta, source, doc_type,
                                section="(语义段)", kind=KIND_TEXT))
    final: list[Chunk] = []
    for c in chunks:
        for piece in hard_cap(c.text, max_chars, overlap):
            final.append(_mk_chunk(piece, meta, source, doc_type,
                                   section=c.meta["section"], kind=KIND_TEXT))
    for b in blocks:  # 表格与结构化原子块独立入库
        if b.kind == KIND_TABLE or b.atomic:
            for piece in hard_cap(b.text, max(max_chars, 1200), overlap):
                final.append(_mk_chunk(piece, meta, source, doc_type,
                                       section="(表格)" if b.kind == KIND_TABLE else "(活动规则)",
                                       kind=b.kind, page=b.page, sheet=b.sheet,
                                       valid_from=b.valid_from, valid_until=b.valid_until))
    return final


def apply_parent_child(chunks: list[Chunk], parent_max: int = 800,
                       child_max: int = 400, overlap: int = 60) -> list[Chunk]:
    """parent-child：超长 chunk 拆成子块（检索精度高），父节全文挂 meta.parent_text（返回上下文完整）。

    短 chunk 原样保留（本身就是自己的父节）。生产建议 parent_text 外挂文档库，此处随元数据走便于演示。
    """
    out: list[Chunk] = []
    for c in chunks:
        if len(c.text) <= parent_max:
            out.append(c)
            continue
        for piece in hard_cap(c.text, child_max, overlap):
            child_meta = dict(c.meta)
            child_meta["parent_text"] = c.text
            out.append(Chunk(id="", text=piece, meta=child_meta))
    return out


# ---------------------------------------------------------------------------
# embedder —— 嵌入层（OpenAI 兼容端点，SiliconFlow Qwen3-Embedding-8B）
# ---------------------------------------------------------------------------

class Embedder:
    """批量嵌入 + 失败指数退避重试（SiliconFlow 免费档限流常见）+ 本地磁盘缓存。

    缓存按 sha256(model+text) 键控：语料没变的块重跑直接命中缓存（0 API 调用），
    只有新增/变更的块才真正嵌入——调试期反复重跑不再重复烧配额。
    """

    def __init__(self, api_base: str, api_key: str, model: str, batch_size: int = 16,
                 cache_file: str | None = None):
        try:
            from openai import OpenAI
        except ImportError as e:
            raise SystemExit(f"import openai 失败：{e!r}（pip install openai）") from e
        self.client = OpenAI(api_key=api_key or "sk-placeholder", base_url=api_base)
        self.api_key = api_key  # 惰性校验：全命中缓存时无需 key（_embed_batch_retry 才检查）
        self.model = model
        self.batch_size = batch_size
        self.dimension: int | None = None
        self.cache_file = Path(cache_file) if cache_file else None
        self._cache: dict[str, list[float]] = {}
        self._cache_hits = 0
        if self.cache_file and self.cache_file.exists():
            try:
                self._cache = json.loads(self.cache_file.read_text(encoding="utf-8"))
                print(f"[cache] 载入嵌入缓存 {len(self._cache)} 条（{self.cache_file}）", file=sys.stderr)
            except Exception as e:
                print(f"[cache] 缓存文件损坏，忽略重建：{e}", file=sys.stderr)
                self._cache = {}

    @staticmethod
    def _cache_key(model: str, text: str) -> str:
        return hashlib.sha256(f"{model}\x00{text}".encode("utf-8")).hexdigest()

    def embed(self, texts: list[str]) -> list[list[float]]:
        keys = [self._cache_key(self.model, t[:4000]) for t in texts]
        uncached: list[int] = []          # 既不在磁盘缓存、也不在本批次内重复的索引
        seen_in_batch: set[str] = set()
        for i, k in enumerate(keys):
            if k not in self._cache and k not in seen_in_batch:
                uncached.append(i)
                seen_in_batch.add(k)
        if not uncached:
            print(f"[cache] {len(texts)} 块全部命中缓存，本次 0 次 API 调用", file=sys.stderr)
        else:
            print(f"[cache] 命中 {len(keys) - len(uncached)}，需嵌入 {len(uncached)} 块", file=sys.stderr)
            for i in range(0, len(uncached), self.batch_size):
                batch_idx = uncached[i : i + self.batch_size]
                batch = [texts[j][:4000] for j in batch_idx]
                vectors = self._embed_batch_retry(batch)
                for j, v in zip(batch_idx, vectors):
                    self._cache[keys[j]] = v
                print(f"  [embed] API {min(i + self.batch_size, len(uncached))}/{len(uncached)}", file=sys.stderr)
            self._save_cache()
        result = [self._cache[k] for k in keys]
        if result and self.dimension is None:
            self.dimension = len(result[0])
            print(f"[embed] 维度确认：{self.dimension}（须与 Java 侧一致）", file=sys.stderr)
        return result

    def embed_query(self, text: str) -> list[float]:
        key = self._cache_key(self.model, text[:4000])
        if key in self._cache:
            self._cache_hits += 1
            return self._cache[key]
        vector = self._embed_batch_retry([text])[0]
        self._cache[key] = vector
        self._save_cache()
        return vector

    def _save_cache(self) -> None:
        if not self.cache_file:
            return
        try:
            self.cache_file.write_text(json.dumps(self._cache), encoding="utf-8")
        except Exception as e:
            print(f"[cache] 缓存写入失败（不影响本次结果）：{e}", file=sys.stderr)

    def _embed_batch_retry(self, batch: list[str], retries: int = 5) -> list[list[float]]:
        if not self.api_key:
            raise SystemExit(
                "需要嵌入 API：export SF_KEY=<SiliconFlow key>"
                "（若语料未变应全部命中缓存，不会走到这里；语料有变时首次需配 key）")
        delay = 2.0
        for attempt in range(1, retries + 1):
            try:
                resp = self.client.embeddings.create(model=self.model, input=batch)
                return [d.embedding for d in resp.data]  # data 顺序与输入一致
            except Exception as e:
                if attempt == retries:
                    raise
                print(f"[embed] 批次失败（{e}），{delay:.0f}s 后重试 {attempt}/{retries}", file=sys.stderr)
                time.sleep(delay)
                delay = min(delay * 2, 30)
        raise RuntimeError("unreachable")


# ---------------------------------------------------------------------------
# writers —— 落库层（Chroma / Milvus，确定性 id 幂等 upsert）
# ---------------------------------------------------------------------------

def finalize_ids(chunks: list[Chunk]) -> None:
    """确定性 id：sha256(source#section#text) 前 32 位 —— 重跑幂等覆盖而非重复灌库。"""
    for c in chunks:
        raw = f"{c.meta['source']}#{c.meta['section']}#{c.text}"
        c.id = hashlib.sha256(raw.encode("utf-8")).hexdigest()[:32]


def _preflight_tcp(host: str, port: int, timeout: float = 5.0) -> None:
    """落库前 TCP 预检：向量库不可达时 5 秒内给出可操作的诊断，而不是无限挂起。"""
    import socket

    try:
        with socket.create_connection((host, port), timeout=timeout):
            return
    except OSError as e:
        raise SystemExit(
            f"\n无法连接 Chroma {host}:{port}（{e.__class__.__name__}: {e}）\n"
            f"  ① 云安全组/防火墙是否放行 {port}？——超时=包被丢弃，拒绝=端口没监听\n"
            f"  ② 服务端只绑本机时用 SSH 隧道：ssh -L {port}:localhost:{port} <user>@{host}，"
            f"然后 --uri http://localhost:{port}\n"
            f"  ③ 服务器本地验证 Chroma 是否活着：curl http://localhost:{port}/api/v2/heartbeat")


class ChromaWriter:
    def __init__(self, uri: str, collection: str, token: str | None = None):
        try:
            import chromadb
        except ImportError as e:
            # SystemExit 不会打印异常链，必须把真实 ImportError 原文带出来，否则无法诊断
            raise SystemExit(
                f"\nimport chromadb 失败：{e!r}\n"
                f"  pip 显示已装但 import 失败 → 多为安装损坏或解释器不一致：\n"
                f"  ① python -c \"import chromadb\"  看真实报错\n"
                f"  ② which -a python pip           确认 pip 装的与 python 跑的是同一个环境\n"
                f"  ③ pip install --force-reinstall --no-cache-dir chromadb") from e
        scheme, host, port = _parse_uri(uri, default_port=8000)
        _preflight_tcp(host, port)  # 先 5s 预检，避免 chromadb 客户端对不可达端口无限挂起
        # 远程服务器支持：https + Bearer token 鉴权（服务端开启 TokenAuthentication 时必传）
        headers = {"Authorization": f"Bearer {token}"} if token else None
        self.client = chromadb.HttpClient(host=host, port=port,
                                          ssl=(scheme == "https"), headers=headers)
        self.collection = self.client.get_or_create_collection(
            name=collection, metadata={"hnsw:space": "cosine"})

    def upsert(self, chunks: list[Chunk], vectors: list[list[float]]) -> None:
        bs = 128
        for i in range(0, len(chunks), bs):
            part, vecs = chunks[i : i + bs], vectors[i : i + bs]
            self.collection.upsert(
                ids=[c.id for c in part],
                embeddings=vecs,
                documents=[c.text for c in part],
                metadatas=[c.meta for c in part],
            )

    def count(self) -> int:
        return self.collection.count()

    def recreate(self) -> None:
        """删除并重建 collection（--rebuild：语料是唯一事实源，库是纯派生物，删减语料后
        upsert 不清旧块，必须整库重建才不留孤儿数据）。"""
        self.client.delete_collection(self.collection.name)
        self.collection = self.client.get_or_create_collection(
            name=self.collection.name, metadata={"hnsw:space": "cosine"})

    def search(self, vector: list[float], top_k: int = 5) -> list[dict]:
        res = self.collection.query(query_embeddings=[vector], n_results=top_k)
        return [{"text": res["documents"][0][i],
                 "meta": res["metadatas"][0][i],
                 "score": res["distances"][0][i]}  # chroma 返回距离，越小越近
                for i in range(len(res["ids"][0]))]


def _milvus_dim(desc: dict) -> int | None:
    for f in desc.get("fields", []):
        if "vector" in str(f.get("name", "")):
            dim = (f.get("params") or {}).get("dim")
            return int(dim) if dim else None
    return None


class MilvusWriter:
    def __init__(self, uri: str, collection: str, dimension: int):
        try:
            from pymilvus import MilvusClient
        except ImportError as e:
            raise SystemExit(f"import pymilvus 失败：{e!r}（pip install pymilvus）") from e
        self.client = MilvusClient(uri=uri)
        self.collection = collection
        self.dimension = dimension
        if not self.client.has_collection(collection):
            self.client.create_collection(
                collection_name=collection,
                dimension=dimension,
                metric_type="COSINE",
                auto_id=False,
                id_type="string",
                primary_field_name="id",
                vector_field_name="vector",
            )
        else:
            dim = _milvus_dim(self.client.describe_collection(collection))
            if dim is not None and dim != dimension:
                raise SystemExit(
                    f"维度不匹配：Milvus 中已有 collection「{collection}」维度 {dim}，"
                    f"当前嵌入维度 {dimension}。换嵌入模型必须全库重灌（drop collection 后重跑）。")

    def upsert(self, chunks: list[Chunk], vectors: list[list[float]]) -> None:
        bs = 128
        for i in range(0, len(chunks), bs):
            data = [{"id": c.id, "vector": v, "text": c.text, **c.meta}
                    for c, v in zip(chunks[i : i + bs], vectors[i : i + bs])]
            self.client.upsert(collection_name=self.collection, data=data)

    def count(self) -> int:
        return int(self.client.get_collection_stats(self.collection).get("row_count", 0))

    def recreate(self) -> None:
        """--rebuild：drop 后按原维度重建（清语料删减后的孤儿块）。"""
        self.client.drop_collection(self.collection)
        self.client.create_collection(
            collection_name=self.collection,
            dimension=self.dimension,
            metric_type="COSINE",
            auto_id=False,
            id_type="string",
            primary_field_name="id",
            vector_field_name="vector",
        )

    def search(self, vector: list[float], top_k: int = 5) -> list[dict]:
        res = self.client.search(
            collection_name=self.collection, data=[vector], limit=top_k,
            output_fields=["text", "source", "section", "doc_type", "temporal_tag"])
        return [{"text": item.get("entity", {}).get("text", ""),
                 "meta": item.get("entity", {}),
                 "score": item.get("distance")}
                for item in (res[0] if res else [])]


def _parse_uri(uri: str, default_port: int) -> tuple[str, str, int]:
    """解析 http(s)://host[:port] → (scheme, host, port)。"""
    m = re.match(r"^(https?)://([^/:]+)(?::(\d+))?$", uri)
    if not m:
        raise SystemExit(f"无法解析 uri（须为 http://host:port 或 https://host:port）：{uri}")
    return m.group(1), m.group(2), int(m.group(3) or default_port)


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------

VERIFY_QUESTIONS = [
    "退款多久到账",
    "七天无理由退货从哪天开始算",
    "增值税专用发票需要提供什么信息",
    "满多少元包邮",
    "会员等级怎么升级",
    "现在有什么满减活动",
    "双 11 活动什么时候开始",
]


def collect_files(input_path: Path) -> list[Path]:
    files = sorted(p for p in input_path.rglob("*")
                   if p.is_file() and p.suffix.lower() in READERS)
    if not files:
        raise SystemExit(f"目录中没有可识别的文档（支持 {', '.join(SUPPORTED_EXTS)}）：{input_path}")
    return files


def run(args: argparse.Namespace) -> None:
    input_path = Path(args.input)
    files = collect_files(input_path if input_path.is_dir() else input_path.parent)
    print(f"[scan] 待处理文档 {len(files)} 个（目录：{input_path}）")

    embedder: Embedder | None = None
    all_chunks: list[Chunk] = []
    for f in files:
        suffix = f.suffix.lower()
        blocks, meta = READERS[suffix](f)
        source = str(f.relative_to(input_path) if input_path.is_dir() else f.name)
        for b in blocks:
            b.source = source
        blocks = clean_document(blocks)
        if not blocks:
            print(f"[skip] {source}：清洗后无有效内容", file=sys.stderr)
            continue
        doc_type = DOC_TYPE_LABEL.get(suffix.lstrip("."), suffix.lstrip("."))
        if args.chunker == "semantic":
            if embedder is None:
                embedder = get_embedder(args)
            chunks = semantic_chunks(blocks, meta, source, doc_type, embedder,
                                     args.semantic_threshold, args.max_chars, args.overlap)
        else:
            chunks = structure_chunks(blocks, meta, source, doc_type,
                                      args.max_chars, args.overlap)
        if args.parent_child:
            chunks = apply_parent_child(chunks)
        finalize_ids(chunks)
        all_chunks.extend(chunks)
        kinds: dict[str, int] = {}
        for c in chunks:
            kinds[c.meta["kind"]] = kinds.get(c.meta["kind"], 0) + 1
        print(f"[parse] {(meta.doc_no or source):<24} type={doc_type:<8} "
              f"temporal={meta.temporal_tag:<10} chunks={len(chunks):<4} {kinds}")

    n_table = sum(1 for c in all_chunks if c.meta["kind"] == KIND_TABLE)
    n_hist = sum(1 for c in all_chunks if c.meta.get("temporal_tag") == "HISTORICAL")
    by_type: dict[str, int] = {}
    for c in all_chunks:
        by_type[c.meta["doc_type"]] = by_type.get(c.meta["doc_type"], 0) + 1
    print(f"\n[chunk] 共 {len(all_chunks)} 块（表格原子块 {n_table}，历史块 {n_hist}）")
    print(f"[chunk] 按文档类型：{by_type}")

    if args.dry_run:
        print("\n[dry-run] 仅解析切分验证，未落库。抽查前 3 块：")
        for c in all_chunks[:3]:
            print(f"  --- {c.meta['source']} · {c.meta['section']} · kind={c.meta['kind']} ---")
            print("  " + c.text[:120].replace("\n", "\n  "))
        return

    if embedder is None:
        embedder = get_embedder(args)
    print(f"\n[embed] 开始嵌入 {len(all_chunks)} 块（model={args.embed_model}, batch={args.batch_size}）")
    vectors = embedder.embed([c.text for c in all_chunks])
    dimension = len(vectors[0])

    print(f"[write] 目标库：{args.store} @ {args.uri} / collection={args.collection}")
    if args.store == "chroma":
        token = args.chroma_token or os.environ.get("CHROMA_TOKEN") or None
        writer = ChromaWriter(args.uri, args.collection, token=token)
    else:
        writer = MilvusWriter(args.uri, args.collection, dimension)
    if args.rebuild:
        print("[write] --rebuild：清空重建 collection（清除语料删减后 upsert 留下的孤儿块）")
        writer.recreate()
    writer.upsert(all_chunks, vectors)
    print(f"[write] 落库完成，库内总量：{writer.count()} 块")

    if not args.skip_verify:
        print("\n[verify] 语义召回自检（top1）：")
        for q in VERIFY_QUESTIONS:
            hits = writer.search(embedder.embed_query(q), top_k=1)
            if hits:
                top, m = hits[0], hits[0]["meta"]
                print(f"  Q：{q}")
                print(f"    → [{m.get('source', '?')} · {m.get('section', '?')}] {top['text'][:60]}…")


def get_embedder(args: argparse.Namespace) -> Embedder:
    # key 惰性校验：语料未变时全命中缓存，不配 key 也能跑（真要调 API 时才检查）
    api_key = args.embed_api_key or os.environ.get(args.embed_api_key_env, "")
    return Embedder(args.embed_api_base, api_key, args.embed_model, args.batch_size,
                    cache_file=None if args.no_cache else args.cache_file)


def build_argparser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="多格式文档清洗-切分-入库流水线（客服知识库 RAG 前置）",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    parser.add_argument("--input", required=True, help="语料目录或单文件路径")
    parser.add_argument("--store", choices=["chroma", "milvus"], default="chroma", help="向量库类型")
    parser.add_argument("--uri", default=None,
                        help="向量库地址（chroma 默认 http://localhost:8000，milvus 默认 http://localhost:19530）")
    parser.add_argument("--chroma-token", default=None,
                        help="Chroma 服务端鉴权 token（Bearer），也可用环境变量 CHROMA_TOKEN；服务端未开鉴权则不用传")
    parser.add_argument("--collection", default="kb_customer_service", help="collection 名称")
    parser.add_argument("--chunker", choices=["structure", "semantic"], default="structure",
                        help="structure=按标题层级（有结构文档，默认）；semantic=句子嵌入断点（txt/扫描 OCR 文本）")
    parser.add_argument("--parent-child", action="store_true",
                        help="parent-child：超长块拆子块检索、父节全文挂 parent_text 元数据")
    parser.add_argument("--max-chars", type=int, default=600, help="单块最大字符数（≈512 token 兜底）")
    parser.add_argument("--overlap", type=int, default=80, help="硬约束切分的重叠字符数")
    parser.add_argument("--semantic-threshold", type=float, default=0.55,
                        help="semantic 切分的相邻句余弦断点阈值（越大越碎）")
    parser.add_argument("--embed-model", default="Qwen/Qwen3-Embedding-8B",
                        help="嵌入模型（必须与 Java 侧同源，换模型=全库重灌）")
    parser.add_argument("--embed-api-base", default="https://api.siliconflow.cn/v1",
                        help="嵌入 API 基地址（OpenAI 兼容）")
    parser.add_argument("--embed-api-key", default=None, help="嵌入 API key（明文，仅测试用）")
    parser.add_argument("--embed-api-key-env", default="SF_KEY",
                        help="嵌入 API key 的环境变量名（推荐方式）")
    parser.add_argument("--batch-size", type=int, default=16, help="嵌入批大小（免费档限流时调小）")
    parser.add_argument("--cache-file", default="embedding-cache.json",
                        help="嵌入磁盘缓存文件（按文本哈希；语料没变的块重跑 0 API 调用）")
    parser.add_argument("--no-cache", action="store_true", help="禁用嵌入缓存（全量真嵌）")
    parser.add_argument("--rebuild", action="store_true",
                        help="落库前清空重建 collection（语料删减后必用，否则旧块残留）")
    parser.add_argument("--dry-run", action="store_true", help="只解析切分打印统计，不嵌入不落库")
    parser.add_argument("--skip-verify", action="store_true", help="跳过落库后的召回自检")
    return parser


def _env_guard() -> None:
    """多 Python 共存防护（真实事故：本机 python=3.9 / python3=3.13 并存，`python` 命令被
    系统 3.9 框架抢占——`which python` 虽指向 venv，shell 命令缓存仍落旧解释器，于是
    pip 装进 3.13 的 venv、跑起来却是 3.9，以 ModuleNotFoundError 收场）。

    入口处校验：当前解释器必须处于脚本旁的 .venv 中（存在时），否则立刻给出修复指引，
    不让用户等到最后一步才见到 import 报错。
    """
    import sys

    print(f"[env] Python {sys.version.split()[0]} @ {sys.executable}", file=sys.stderr)
    in_venv = sys.prefix != getattr(sys, "base_prefix", sys.prefix)
    expected = Path(__file__).resolve().parent / ".venv"
    if not in_venv and expected.exists():
        raise SystemExit(
            f"\n[环境检查] 当前解释器不在虚拟环境中（venv 就在旁边：{expected}），"
            f"依赖都装在 .venv 里，继续跑必然 ModuleNotFoundError。\n"
            f"  修复（二选一）：\n"
            f"    ① 改用 python3 运行：python3 ingest.py ...\n"
            f"       （venv 由 python3 创建；多版本共存时 python 命令可能被系统旧版本抢占）\n"
            f"    ② source {expected}/bin/activate 后清 shell 命令缓存再跑：\n"
            f"       bash: hash -r    zsh: rehash\n"
            f"       然后 python --version 确认是 3.13")


def main() -> None:
    _env_guard()
    args = build_argparser().parse_args()
    if args.uri is None:
        args.uri = "http://localhost:8000" if args.store == "chroma" else "http://localhost:19530"
    run(args)


if __name__ == "__main__":
    main()
